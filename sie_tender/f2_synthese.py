"""F2 (suite) : activités SIE par lot, matrice d'interfaces et synthèse des lots retenus.

Tout part des exigences extraites (phrases réelles du DCE, avec leur source) :
rien n'est reformulé ni ajouté. Les « questions à poser » sont des manques
constatés (champ introuvable, limite de prestation floue...), formulés comme
tels, jamais des réponses supposées.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .models import CategorieExigence, Consultation, Exigence, Lot, Piece, Source, Statut, TypePiece
from .texte import articles, normaliser

# Activités SIE (cahier des charges, § 3).
ACTIVITES: dict[str, list[str]] = {
    "CFO": [r"courants? forts?", r"\bcfo\b", r"eclairage", r"distribution electrique", r"tableaux? (?:electriques?|divisionnaires?|generaux?)",
            r"\belectricite\b(?! courants faibles)", r"\btgbt\b"],
    "CFA": [r"courants? faibles?", r"\bcfa\b", r"appel malade", r"controle d'acces", r"cablage informatique|precablage|\bvdi\b|rj ?45",
            r"interphon", r"videosurveillance|videoprotection", r"intrusion"],
    "SSI": [r"\bssi\b", r"securite incendie", r"detection incendie", r"alarme (?:incendie|de type)", r"\bcmsi\b"],
    "CVC": [r"\bcvc\b", r"chauffage", r"ventilation", r"climatisation", r"traitement d'air"],
    "Plomberie": [r"plomberie", r"sanitaires?"],
    "Désenfumage": [r"desenfumage"],
}

# Nature des interfaces (cahier des charges, F2).
THEMES: dict[str, list[str]] = {
    "Alimentation d'équipements": [r"alimentation", r"\balimente", r"coffret", r"raccordement"],
    "Réservations et cheminements": [r"reservations?", r"percements?", r"cheminements?", r"fourreaux", r"chemins? de cables", r"tremies?"],
    "Sécurité incendie": [r"\bssi\b", r"securite incendie", r"coupe-feu", r"alarme", r"asservi"],
    "Désenfumage": [r"desenfumage"],
    "Limites de prestations": [r"limites? de prestations?", r"a la charge du lot", r"fourni(?:e|s)? (?:et pose(?:e|s)? )?par le lot", r"\bfournit\b"],
    "Travaux préparatoires": [r"travaux preparatoires", r"\bdepose\b", r"demolition", r"provisoire"],
    "Coordination": [r"coordination", r"interfaces? avec", r"en accord avec", r"\bsynthese\b", r"compte prorata", r"tous corps d'etat"],
    "Phasage / site occupé": [r"site occupe", r"milieu occupe", r"\bphasage\b", r"en fonctionnement", r"maintien (?:en activite|de l'activite)"],
}

_MOTS_VIDES = {"lot", "lots", "des", "les", "et", "de", "du", "la", "le", "en", "sur", "pour", "avec", "travaux", "unique"}


def activites(texte: str) -> list[str]:
    n = normaliser(texte)
    return [a for a, motifs in ACTIVITES.items() if any(re.search(m, n) for m in motifs)]


def activites_des_lots(lots: list[Lot], pieces: list[Piece], objet: str | None = None) -> None:
    """Rattache chaque lot aux activités SIE d'après son intitulé et les titres d'articles de son CCTP.

    On ne lit pas le corps du CCTP : un CCTP électricité cite souvent la CVC à propos d'une interface.
    """
    for lot in lots:
        titres = [lot.intitule]
        if len(lots) == 1 and objet:
            titres.append(objet)
        for p in pieces:
            if p.type == TypePiece.CCTP and p.pages and (p.lot == lot.numero or (len(lots) == 1 and not p.lot)):
                titres += [a.titre for a in articles([(pg.numero, pg.texte) for pg in p.pages])]
                titres.append(p.pages[0].texte[:200])
        lot.activites = activites(" ".join(titres))


# --- Matrice d'interfaces ---------------------------------------------------------
@dataclass
class Interface:
    origine: str            # lot dont la pièce parle, ou "commun"
    concernes: list[str]    # lots cités ; ["tous"] pour une contrainte commune
    themes: list[str]
    texte: str
    source: Source
    impacte_retenu: bool = False


def _lots_nommes(n: str, lots: list[Lot]) -> list[str]:
    """Lots cités par leur nom (« le lot CVC ») plutôt que par leur numéro."""
    out = []
    for lot in lots:
        mots = [m for m in re.findall(r"[a-z]{3,}", normaliser(lot.intitule)) if m not in _MOTS_VIDES]
        if any(re.search(rf"\blots?\s+{re.escape(m)}", n) for m in mots):
            out.append(lot.numero)
    return out


def themes(texte: str) -> list[str]:
    n = normaliser(texte)
    return [t for t, motifs in THEMES.items() if any(re.search(m, n) for m in motifs)] or ["Autre"]


def matrice_interfaces(c: Consultation) -> list[Interface]:
    from .f2_exigences import _lots_cites

    retenus = {l.numero for l in c.lots if l.retenu}
    out: list[Interface] = []
    for e in c.exigences:
        commune = e.lot is None
        if e.categorie == CategorieExigence.INTERFACE:
            n = normaliser(e.texte)
            concernes = sorted({x for x in _lots_cites(n) + _lots_nommes(n, c.lots) if x != e.lot}, key=lambda x: x.zfill(3))
            if not concernes:
                concernes = ["tous"] if commune else ["?"]
        elif e.categorie == CategorieExigence.SITE_OCCUPE and commune:
            concernes = ["tous"]
        else:
            continue
        i = Interface(e.lot or "commun", concernes, themes(e.texte), e.texte, e.source)
        i.impacte_retenu = bool(retenus) and (e.lot in retenus or "tous" in concernes or bool(retenus & set(concernes)))
        out.append(i)
    return out


def tableau_croise(interfaces: list[Interface], lots: list[Lot]) -> tuple[list[str], dict[str, dict[str, int]]]:
    """Nombre d'interfaces par couple (lot d'origine, lot concerné)."""
    noms = [l.numero for l in lots]
    axes = noms + [x for x in ("commun",) if any(i.origine == "commun" for i in interfaces)]
    grille = {o: {d: 0 for d in noms} for o in axes}
    for i in interfaces:
        cibles = noms if "tous" in i.concernes else [x for x in i.concernes if x in noms]
        for d in cibles:
            if i.origine in grille and d != i.origine:
                grille[i.origine][d] += 1
    return axes, grille


# --- RSE : heures d'insertion -----------------------------------------------------
_HEURES = re.compile(r"(\d{1,3}(?:[ .  ]\d{3})*|\d+)\s*(?:heures?|h)\b", re.IGNORECASE)


def heures_insertion(exigences: list[Exigence]) -> list[tuple[int, Exigence]]:
    """Nombre d'heures d'insertion demandé, lu dans la phrase même (jamais estimé)."""
    out = []
    for e in exigences:
        if e.categorie != CategorieExigence.RSE or e.sous_categorie != "insertion":
            continue
        m = _HEURES.search(e.texte)
        if m:
            out.append((int(re.sub(r"\D", "", m.group(1))), e))
    return out


# --- Synthèse par lot retenu ------------------------------------------------------
@dataclass
class Element:
    texte: str
    source: Source | None = None


@dataclass
class SyntheseLot:
    lot: Lot
    prestations: list[Element] = field(default_factory=list)
    exigences_techniques: list[Element] = field(default_factory=list)
    risques: list[Element] = field(default_factory=list)
    interfaces: list[Interface] = field(default_factory=list)
    rse: list[Element] = field(default_factory=list)
    heures_insertion: list[Element] = field(default_factory=list)
    planning: list[Element] = field(default_factory=list)
    questions: list[str] = field(default_factory=list)


def _du_lot(e: Exigence, lot: str) -> bool:
    return e.lot in (None, lot)


def synthese_lot(c: Consultation, lot: Lot, pieces: list[Piece] | None = None) -> SyntheseLot:
    from .export import LIBELLES

    s = SyntheseLot(lot)
    unique = len(c.lots) == 1
    cctp = [p for p in (pieces or []) if p.type == TypePiece.CCTP and p.pages and (p.lot == lot.numero or (unique and not p.lot))]
    for p in cctp:
        for a in articles([(pg.numero, pg.texte) for pg in p.pages]):
            if a.titre:
                s.prestations.append(Element(f"Art. {a.numero} : {a.titre}", Source(piece_id=p.id, fichier=p.fichier, page=a.page, article=a.numero)))
    ex = [e for e in c.exigences if _du_lot(e, lot.numero)]
    el = lambda e: Element(e.texte, e.source)  # noqa: E731
    s.exigences_techniques = [el(e) for e in ex if e.categorie == CategorieExigence.TECHNIQUE]
    s.risques = [el(e) for e in ex if e.categorie in (CategorieExigence.ELIMINATOIRE, CategorieExigence.PENALITE, CategorieExigence.SITE_OCCUPE)]
    s.rse = [Element(f"[{e.sous_categorie}] {e.texte}", e.source) for e in ex if e.categorie == CategorieExigence.RSE]
    s.heures_insertion = [Element(f"{h} heures", e.source) for h, e in heures_insertion(ex)]
    s.planning = [el(e) for e in ex if e.categorie == CategorieExigence.PLANNING]
    s.interfaces = [i for i in matrice_interfaces(c)
                    if i.origine == lot.numero or lot.numero in i.concernes or "tous" in i.concernes]

    # Questions : uniquement des manques constatés.
    q = s.questions
    limite_q = c.fiche.date_limite_questions.valeur
    for k, ch in c.fiche.champs().items():
        if ch.statut == Statut.A_VERIFIER and k not in ("intervenants",):
            q.append(f"{LIBELLES.get(k, k)} : non trouvé dans le DCE, à confirmer"
                     + (f" (questions à l'acheteur : {limite_q})" if limite_q else "") + ".")
    if pieces is not None and not cctp:
        q.append(f"Aucun CCTP propre au lot {lot.numero} identifié : vérifier les pièces techniques.")
    for i in s.interfaces:
        if "?" in i.concernes or ("Limites de prestations" in i.themes and i.origine != lot.numero):
            q.append(f"Limite de prestation à confirmer : « {i.texte[:220]} » ({i.source.fichier}, p. {i.source.page}).")
    total = sum(cr.ponderation or 0 for cr in c.criteres)
    if not c.criteres:
        q.append("Critères de notation non trouvés dans le RC.")
    elif abs(total - 100) > 0.5:
        q.append(f"La somme des pondérations fait {total:g} % : critères à confirmer.")
    if c.fiche.visite_obligatoire.valeur == "oui" and not c.fiche.responsable_visite:
        q.append("Visite obligatoire : désigner le collaborateur chargé de la visite.")
    if not s.heures_insertion and any(e.sous_categorie == "insertion" for e in ex if e.categorie == CategorieExigence.RSE):
        q.append("Clause d'insertion sans nombre d'heures lisible : à vérifier.")
    return s


def syntheses_retenues(c: Consultation, pieces: list[Piece] | None = None) -> list[SyntheseLot]:
    return [synthese_lot(c, l, pieces) for l in c.lots if l.retenu]


def _ref(src: Source | None) -> str:
    from .export import ref
    return ref(src)


def exporter_synthese_docx(c: Consultation, pieces: list[Piece] | None, chemin: Path) -> None:
    import docx
    from docx.shared import Pt, RGBColor

    d = docx.Document()
    titre = c.fiche.intitule.valeur or c.fiche.objet.valeur or "consultation"
    d.add_heading(f"Synthèse des lots retenus : {titre}", 0)
    p = d.add_paragraph("Chaque point reprend une phrase du DCE avec sa source. Les questions listent des informations "
                        "absentes ou floues : elles ne présument pas de la réponse.")
    p.runs[0].italic = True

    def puces(titre_section: str, elements: list[Element], vide: str = "Rien de trouvé dans le DCE.") -> None:
        d.add_heading(titre_section, 2)
        if not elements:
            d.add_paragraph(vide).runs[0].italic = True
        for e in elements:
            par = d.add_paragraph(e.texte, style="List Bullet")
            if e.source:
                r = par.add_run(f"  [{_ref(e.source)}]")
                r.font.size, r.italic = Pt(8), True

    syntheses = syntheses_retenues(c, pieces)
    if not syntheses:
        d.add_paragraph("Aucun lot retenu : la synthèse est produite après la décision GO et la sélection des lots.")
    if c.criteres:
        d.add_heading("Critères de notation (RC)", 1)
        for cr in c.criteres:
            d.add_paragraph(f"{cr.intitule} : {cr.ponderation:g} %  [{_ref(cr.source)}]", style="List Bullet")
            for sc in cr.sous_criteres:
                d.add_paragraph(f"{sc.intitule} : {sc.ponderation:g} %", style="List Bullet 2")
    for s in syntheses:
        acts = f" ({', '.join(s.lot.activites)})" if s.lot.activites else ""
        d.add_heading(f"Lot {s.lot.numero} : {s.lot.intitule}{acts}", 1)
        puces("Prestations (articles du CCTP)", s.prestations)
        puces("Exigences techniques", s.exigences_techniques)
        puces("Risques : clauses éliminatoires, pénalités, site occupé", s.risques)
        puces("Interfaces et impacts des autres lots",
              [Element(f"{i.origine} → {', '.join(i.concernes)} · {', '.join(i.themes)} : {i.texte}", i.source) for i in s.interfaces])
        puces("RSE", s.rse)
        if s.heures_insertion:
            puces("Heures d'insertion demandées", s.heures_insertion)
        puces("Planning", s.planning)
        d.add_heading("Questions et points à vérifier", 2)
        for q in s.questions or ["Aucun manque relevé."]:
            r = d.add_paragraph(style="List Bullet").add_run(q)
            r.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)
    d.save(chemin)

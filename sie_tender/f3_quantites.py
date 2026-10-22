"""F3 : import des quantités relevées par le bureau d'études (fichier Excel SIE).

Aucune quantité n'est calculée à partir des plans : l'outil lit le tableau,
signale les lignes douteuses et laisse le bureau d'études corriger.
"""
from __future__ import annotations

import re
from pathlib import Path

from .models import LigneQuantite, SourceTableur, Statut
from .texte import normaliser

ENTETES = {
    "code": [r"^n[°o]?$", r"^code", r"^ref", r"^article", r"^poste", r"^item", r"^n[°o] ?(?:de )?(?:prix|poste)"],
    "designation": [r"designation", r"libelle", r"description", r"intitule", r"nature des"],
    "unite": [r"^u$", r"^unite", r"^unit"],
    "quantite": [r"^q(?:te|uantite)", r"^qt", r"^nombre", r"^nb"],
    "famille": [r"famille", r"chapitre", r"categorie", r"rubrique"],
    "lot": [r"^lot"],
}

# Familles par mots-clés, quand le tableau n'a pas de colonne « famille ».
FAMILLES = [
    ("Tableaux et protections", r"tableau|tgbt|disjoncteur|interrupteur differentiel|coffret|armoire|parafoudre"),
    ("Câbles et conduits", r"cable|cablage|conduit|goulotte|chemin de cable|icta|u1000|r2v|fil "),
    ("Éclairage", r"luminaire|eclairage|led|hublot|projecteur|spot|bloc autonome|baes|baeh|detecteur de presence"),
    ("Appareillage", r"prise|interrupteur|va-et-vient|bouton|appareillage"),
    ("Sécurité incendie", r"\bssi\b|alarme|detecteur (?:optique|de fumee)|declencheur|desenfumage|sirene|centrale incendie|\bdad\b|\bcmsi\b"),
    ("Courants faibles", r"\bvdi\b|rj45|baie|fibre|telephon|interphon|controle d'acces|video|wifi|sonorisation"),
    ("CVC", r"ventilat|\bcta\b|climatis|radiateur|chaudiere|pompe a chaleur|gaine|vmc|clapet"),
    ("Plomberie", r"plomberie|sanitaire|lavabo|wc|robinet|tube cuivre|per |evacuation"),
    ("Main-d'œuvre et prestations", r"main d'oeuvre|depose|essais|mise en service|doe|etude|forfait|nettoyage"),
]


def famille_de(designation: str) -> str:
    n = normaliser(designation)
    for nom, motif in FAMILLES:
        if re.search(motif, n):
            return nom
    return "Divers"


def _nombre(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _colonnes(ligne: list) -> dict[str, int]:
    cols: dict[str, int] = {}
    for i, v in enumerate(ligne):
        n = normaliser(str(v or ""))
        if not n:
            continue
        for cle, motifs in ENTETES.items():
            if cle not in cols and any(re.search(m, n) for m in motifs):
                cols[cle] = i
                break
    return cols


def importer_quantites(chemin: str | Path, lot_defaut: str | None = None) -> tuple[list[LigneQuantite], list[str]]:
    """Lit toutes les feuilles d'un classeur ; renvoie les lignes et un journal."""
    from openpyxl import load_workbook

    chemin = Path(chemin)
    journal: list[str] = []
    lignes: list[LigneQuantite] = []
    wb = load_workbook(chemin, data_only=True, read_only=True)
    for ws in wb.worksheets:
        rangees = list(ws.iter_rows(values_only=True))
        cols, debut = {}, None
        for i, r in enumerate(rangees[:30]):
            c = _colonnes(list(r))
            if "designation" in c and "quantite" in c:
                cols, debut = c, i + 1
                break
        if debut is None:
            journal.append(f"F3 : feuille « {ws.title} » ignorée (pas d'en-tête Désignation / Quantité).")
            continue
        m = re.search(r"lot\s*(\d{1,2})", normaliser(ws.title))
        lot_feuille = str(int(m.group(1))) if m else lot_defaut
        famille_courante = None
        for j, r in enumerate(rangees[debut:], start=debut + 1):
            r = list(r) + [None] * 10
            des = r[cols["designation"]]
            if des is None or not str(des).strip():
                continue
            des = re.sub(r"\s+", " ", str(des)).strip()
            qte_brute = r[cols["quantite"]]
            qte = _nombre(qte_brute)
            unite = str(r[cols["unite"]]).strip() if "unite" in cols and r[cols["unite"]] not in (None, "") else None
            # Ligne de chapitre : texte sans quantité ni unité
            if qte is None and unite is None and qte_brute in (None, ""):
                famille_courante = des
                continue
            l = LigneQuantite(
                lot=str(r[cols["lot"]]).strip() if "lot" in cols and r[cols["lot"]] else lot_feuille,
                code=str(r[cols["code"]]).strip() if "code" in cols and r[cols["code"]] not in (None, "") else None,
                famille=(str(r[cols["famille"]]).strip() if "famille" in cols and r[cols["famille"]] else None) or famille_courante or famille_de(des),
                designation=des, unite=unite, quantite=qte,
                source=SourceTableur(fichier=chemin.name, feuille=ws.title, ligne=j),
            )
            problemes = []
            if qte is None:
                problemes.append(f"quantité illisible « {qte_brute} »")
            elif qte < 0:
                problemes.append("quantité négative")
            elif qte == 0:
                problemes.append("quantité nulle")
            if unite is None:
                problemes.append("unité manquante")
            if problemes:
                l.statut = Statut.A_VERIFIER
                l.remarque = ", ".join(problemes)
            lignes.append(l)
    a_verif = sum(1 for l in lignes if l.statut == Statut.A_VERIFIER)
    journal.append(f"F3 : {len(lignes)} lignes importées depuis {chemin.name}, dont {a_verif} à vérifier.")
    return lignes, journal

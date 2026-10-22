"""Import d'un DCE : lecture des fichiers, OCR si besoin, classement des pièces."""
from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from .models import Page, Piece, TypePiece
from .texte import nettoyer, normaliser

EXTENSIONS_TEXTE = {".pdf", ".docx", ".txt", ".md", ".xlsx", ".xlsm"}
EXTENSIONS_PLANS = {".dwg", ".dxf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}

# Mots repérés dans le nom de fichier puis, à défaut, dans la première page.
_INDICES_TYPE: list[tuple[TypePiece, list[str]]] = [
    (TypePiece.RC, [r"\brc\b", "reglement de (la )?consultation"]),
    (TypePiece.CCAP, [r"\bccap\b", "cahier des clauses administratives particulieres"]),
    (TypePiece.CCTP, [r"\bcctp\b", "cahier des clauses techniques particulieres"]),
    (TypePiece.AE, [r"\bae\b", "acte d'engagement"]),
    (TypePiece.DPGF, [r"\bdpgf\b", "decomposition du prix global"]),
    (TypePiece.BPU, [r"\bbpu\b", r"\bdqe\b", "bordereau des prix"]),
    (TypePiece.PLAN, [r"\bplan(s)?\b", "synoptique"]),
]

_LOT = re.compile(r"\blot\s*(?:n\s*[°o]?\s*)?(\d{1,2})\b", re.IGNORECASE)


def type_de_piece(nom_fichier: str, premiere_page: str = "") -> TypePiece:
    nom = normaliser(Path(nom_fichier).stem.replace("_", " ").replace("-", " "))
    for t, motifs in _INDICES_TYPE:
        if any(re.search(m, nom) for m in motifs):
            return t
    # Dans le contenu, on retient l'intitulé complet qui apparaît le plus tôt
    # (un CCTP cite souvent le RC plus bas dans sa page de garde).
    debut = normaliser(premiere_page[:1500])
    meilleur: tuple[int, TypePiece] | None = None
    for t, motifs in _INDICES_TYPE:
        for m in motifs[1:]:
            r = re.search(m, debut)
            if r and (meilleur is None or r.start() < meilleur[0]):
                meilleur = (r.start(), t)
    if meilleur:
        return meilleur[1]
    return TypePiece.ANNEXE if debut else TypePiece.AUTRE


def lot_de_piece(nom_fichier: str, type_piece: TypePiece) -> str | None:
    """Numéro de lot d'une pièce d'après son nom (CCTP_lot03.pdf -> "3")."""
    if type_piece in (TypePiece.RC, TypePiece.CCAP):
        return None
    m = _LOT.search(Path(nom_fichier).stem.replace("_", " ").replace("-", " "))
    return str(int(m.group(1))) if m else None


def _lire_pdf(chemin: Path) -> tuple[list[Page], str | None]:
    import pymupdf as fitz

    pages: list[Page] = []
    remarque = None
    with fitz.open(chemin) as doc:
        for i, page in enumerate(doc, start=1):
            texte = page.get_text("text")
            ocr = False
            if len(texte.strip()) < 20 and page.get_images():
                texte_ocr = _ocr_page(page)
                if texte_ocr is None:
                    remarque = "Pages scannées sans texte : OCR non disponible (installer Tesseract)."
                else:
                    texte, ocr = texte_ocr, True
            pages.append(Page(numero=i, texte=nettoyer(texte), ocr=ocr))
    return pages, remarque


def _ocr_page(page) -> str | None:
    """OCR via Tesseract si installé (paquet tesseract-ocr + langue fra)."""
    if shutil.which("tesseract") is None:
        return None
    try:
        tp = page.get_textpage_ocr(language="fra", dpi=300, full=True)
        return page.get_text("text", textpage=tp)
    except Exception:
        return None


def _lire_docx(chemin: Path) -> list[Page]:
    import docx

    d = docx.Document(str(chemin))
    lignes = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lignes.append(" | ".join(c.text.strip() for c in row.cells))
    # Word n'a pas de pages fiables : on découpe par blocs de 60 lignes.
    pages = []
    for i in range(0, max(len(lignes), 1), 60):
        pages.append(Page(numero=i // 60 + 1, texte=nettoyer("\n".join(lignes[i:i + 60]))))
    return pages


def _lire_xlsx(chemin: Path) -> list[Page]:
    from openpyxl import load_workbook

    wb = load_workbook(chemin, read_only=True, data_only=True)
    pages = []
    for i, ws in enumerate(wb.worksheets, start=1):
        lignes = []
        for row in ws.iter_rows(values_only=True):
            cellules = [str(c) for c in row if c is not None]
            if cellules:
                lignes.append(" | ".join(cellules))
        pages.append(Page(numero=i, texte=f"[Feuille {ws.title}]\n" + "\n".join(lignes)))
    return pages


def lire_fichier(chemin: Path) -> Piece:
    ext = chemin.suffix.lower()
    piece = Piece(fichier=chemin.name)
    if ext in EXTENSIONS_PLANS:
        piece.type = TypePiece.PLAN
        piece.lisible = False
        piece.remarque = "Plan stocké sans lecture (métrés faits par le bureau d'études)."
        piece.lot = lot_de_piece(chemin.name, piece.type)
        return piece
    if ext == ".pdf":
        piece.pages, piece.remarque = _lire_pdf(chemin)
    elif ext == ".docx":
        piece.pages = _lire_docx(chemin)
    elif ext in (".xlsx", ".xlsm"):
        piece.pages = _lire_xlsx(chemin)
    elif ext in (".txt", ".md"):
        piece.pages = [Page(numero=1, texte=nettoyer(chemin.read_text(encoding="utf-8", errors="replace")))]
    else:
        piece.lisible = False
        piece.remarque = f"Format {ext or 'inconnu'} non lu."
    piece.nb_pages = len(piece.pages)
    premiere = piece.pages[0].texte if piece.pages else ""
    piece.type = type_de_piece(chemin.name, premiere)
    if piece.type == TypePiece.PLAN and ext == ".pdf":
        # Les plans PDF n'ont pas vocation à être lus : on les garde mais on ne les analyse pas.
        piece.lisible = False
        piece.remarque = "Plan stocké sans lecture."
    if piece.lisible and not any(p.texte.strip() for p in piece.pages):
        piece.lisible = False
        piece.remarque = piece.remarque or "Aucun texte extrait."
    piece.lot = lot_de_piece(chemin.name, piece.type)
    return piece


def importer_dce(chemin: str | Path) -> list[Piece]:
    """Lit un dossier ou une archive ZIP de DCE et renvoie ses pièces classées."""
    chemin = Path(chemin)
    tmp = None
    if chemin.is_file() and chemin.suffix.lower() == ".zip":
        tmp = Path(tempfile.mkdtemp(prefix="dce_"))
        with zipfile.ZipFile(chemin) as z:
            z.extractall(tmp)
        racine = tmp
    elif chemin.is_dir():
        racine = chemin
    elif chemin.is_file():
        return [lire_fichier(chemin)]
    else:
        raise FileNotFoundError(chemin)
    try:
        fichiers = sorted(p for p in racine.rglob("*") if p.is_file() and not p.name.startswith((".", "~$")))
        return [lire_fichier(f) for f in fichiers]
    finally:
        if tmp is not None:
            shutil.rmtree(tmp, ignore_errors=True)

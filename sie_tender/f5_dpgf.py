"""F5 : devis fournisseurs, association aux postes, DPGF provisoire et contrôles.

L'outil associe des prix d'achat aux postes et calcule un prix de vente
uniquement avec les coefficients saisis par SIE (taux horaire, frais, aléas,
marge). Tant qu'un coefficient manque, le prix de vente reste vide.
"""
from __future__ import annotations

import csv
import re
from difflib import SequenceMatcher
from pathlib import Path

from .f3_quantites import _colonnes, _nombre
from .f6_archives import jetons
from .models import Consultation, Constat, Devis, LigneDevis, PosteDPGF
from .texte import normaliser

UNITES_EQUIV = {"u": "u", "un": "u", "unite": "u", "pce": "u", "piece": "u", "ens": "ens", "ensemble": "ens",
                "ft": "ft", "forfait": "ft", "f": "ft", "ml": "ml", "m": "ml", "mlin": "ml", "m2": "m2", "m²": "m2",
                "h": "h", "heure": "h", "kg": "kg", "l": "l"}

ENTETES_PRIX = [r"^p\.?u", r"prix unitaire", r"^prix", r"tarif"]


def unite_norm(u: str | None) -> str | None:
    if not u:
        return None
    n = normaliser(str(u)).replace(".", "").strip()
    return UNITES_EQUIV.get(n, n)


# --- Lecture des devis ------------------------------------------------------------
def _lire_tableur(chemin: Path) -> list[LigneDevis]:
    from openpyxl import load_workbook

    wb = load_workbook(chemin, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rangees = [list(r) + [None] * 10 for r in ws.iter_rows(values_only=True)]
        out += _lignes_depuis_rangees(rangees)
    return out


def _lignes_depuis_rangees(rangees: list[list]) -> list[LigneDevis]:
    for i, r in enumerate(rangees[:30]):
        cols = _colonnes(r)
        prix = next((k for k, v in enumerate(r) if any(re.search(m, normaliser(str(v or ""))) for m in ENTETES_PRIX)), None)
        if "designation" in cols and prix is not None:
            out = []
            for r2 in rangees[i + 1:]:
                des = r2[cols["designation"]]
                pu = _nombre(r2[prix])
                if not des or pu is None:
                    continue
                out.append(LigneDevis(designation=re.sub(r"\s+", " ", str(des)).strip(),
                                      unite=str(r2[cols["unite"]]) if "unite" in cols and r2[cols["unite"]] else None,
                                      quantite=_nombre(r2[cols["quantite"]]) if "quantite" in cols else None,
                                      prix_unitaire=pu))
            return out
    return []


_LIGNE_PDF = re.compile(
    r"^(?P<des>[A-Za-zÀ-ÿ0-9][^\n]{3,120}?)\s+(?P<qte>\d+(?:[.,]\d+)?)\s+(?P<u>u|un|ml|m|m2|m²|ens|ft|forfait|h|kg|pce)\s+"
    r"(?P<pu>\d{1,3}(?:[  .]?\d{3})*(?:,\d{1,2}|\.\d{2})?)\s*(?:€|eur)?(?:\s+\d[\d  .,]*\s*(?:€|eur)?)?\s*$",
    re.IGNORECASE | re.MULTILINE)


def _prix(s: str) -> float | None:
    s = s.replace(" ", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _lire_pdf(chemin: Path) -> list[LigneDevis]:
    import pymupdf

    out = []
    with pymupdf.open(chemin) as doc:
        for page in doc:
            for m in _LIGNE_PDF.finditer(page.get_text("text")):
                out.append(LigneDevis(designation=m.group("des").strip(), unite=m.group("u"),
                                      quantite=_nombre(m.group("qte")), prix_unitaire=_prix(m.group("pu"))))
    return out


def importer_devis(chemin: str | Path, fournisseur: str | None = None, lot: str | None = None) -> Devis:
    chemin = Path(chemin)
    ext = chemin.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        lignes = _lire_tableur(chemin)
    elif ext == ".csv":
        with open(chemin, encoding="utf-8-sig", newline="") as f:
            dialecte = csv.Sniffer().sniff(f.read(2048), delimiters=";,\t")
            f.seek(0)
            lignes = _lignes_depuis_rangees([list(r) + [None] * 10 for r in csv.reader(f, dialecte)])
    elif ext == ".pdf":
        lignes = _lire_pdf(chemin)
    else:
        raise ValueError(f"Format de devis non pris en charge : {ext}")
    return Devis(fournisseur=fournisseur or chemin.stem, fichier=chemin.name, lot=lot, lignes=lignes)


# --- Association devis / postes ------------------------------------------------------
def similarite(a: str, b: str) -> float:
    ja, jb = set(jetons(a)), set(jetons(b))
    jac = len(ja & jb) / len(ja | jb) if ja and jb else 0.0
    seq = SequenceMatcher(None, normaliser(a), normaliser(b)).ratio()
    return round(0.6 * jac + 0.4 * seq, 3)


def associer(c: Consultation, devis: Devis, seuil: float = 0.35) -> None:
    """Propose pour chaque ligne de devis le poste de quantités le plus proche.

    Les associations sous le seuil restent vides ; toutes restent à valider.
    """
    postes = [q for q in c.quantites if devis.lot is None or q.lot in (devis.lot, None)]
    for l in devis.lignes:
        if l.association_validee:
            continue
        meilleur, score = None, 0.0
        for q in postes:
            s = similarite(l.designation, q.designation)
            if unite_norm(l.unite) and unite_norm(q.unite) and unite_norm(l.unite) == unite_norm(q.unite):
                s += 0.05
            if s > score:
                meilleur, score = q, s
        l.ligne_quantite_id = meilleur.id if meilleur and score >= seuil else None
        l.score_association = round(score, 3)


# --- DPGF ----------------------------------------------------------------------------
def _prix_vente(prix_materiel: float | None, temps_h: float | None, c: Consultation) -> float | None:
    k = c.coefficients
    if prix_materiel is None or None in (k.frais_generaux, k.aleas, k.marge):
        return None
    mo = 0.0
    if temps_h:
        if k.taux_horaire is None:
            return None
        mo = temps_h * k.taux_horaire
    debourse = prix_materiel + mo
    return round(debourse * (1 + k.frais_generaux / 100) * (1 + k.aleas / 100) * (1 + k.marge / 100), 2)


def construire_dpgf(c: Consultation, temps_pose: dict[str, float] | None = None) -> list[PosteDPGF]:
    """Un poste par ligne de quantité ; le prix d'achat retenu est le moins cher des devis associés."""
    temps_pose = temps_pose or {p.ligne_quantite_id: p.temps_pose_h for p in c.dpgf if p.temps_pose_h is not None}
    offres: dict[str, list[tuple[float, Devis, LigneDevis]]] = {}
    for d in c.devis:
        for l in d.lignes:
            if l.ligne_quantite_id and l.prix_unitaire is not None:
                offres.setdefault(l.ligne_quantite_id, []).append((l.prix_unitaire, d, l))
    postes = []
    for q in c.quantites:
        p = PosteDPGF(ligne_quantite_id=q.id, lot=q.lot, code=q.code, designation=q.designation, unite=q.unite,
                      quantite=q.quantite, temps_pose_h=temps_pose.get(q.id))
        if q.id in offres:
            pu, d, l = min(offres[q.id], key=lambda x: x[0])
            p.prix_materiel, p.devis_id, p.fournisseur = pu, d.id, d.fournisseur
            if unite_norm(l.unite) and unite_norm(q.unite) and unite_norm(l.unite) != unite_norm(q.unite):
                p.alertes.append(f"Écart d'unité : devis en {l.unite}, métré en {q.unite}")
            if l.quantite is not None and q.quantite is not None and l.quantite + 1e-9 < q.quantite:
                p.alertes.append(f"Quantité du devis ({l.quantite:g}) inférieure au métré ({q.quantite:g})")
            if pu <= 0:
                p.alertes.append("Prix unitaire nul ou négatif")
        elif not (q.famille or "").lower().startswith("main-d"):
            p.alertes.append("Poste sans prix fournisseur")
        if q.quantite is None:
            p.alertes.append("Quantité manquante")
        p.prix_unitaire_vente = _prix_vente(p.prix_materiel, p.temps_pose_h, c)
        if p.prix_unitaire_vente is not None and p.quantite is not None:
            p.total = round(p.prix_unitaire_vente * p.quantite, 2)
        postes.append(p)
    # Prix d'achat anormalement éloigné de la médiane des devis sur le même poste
    for qid, lst in offres.items():
        if len(lst) >= 2:
            prix = sorted(x[0] for x in lst)
            med = prix[len(prix) // 2]
            if med and (prix[-1] > 3 * med or prix[0] < med / 3):
                poste = next(p for p in postes if p.ligne_quantite_id == qid)
                poste.alertes.append(f"Écart important entre devis ({prix[0]:g} à {prix[-1]:g} €)")
    return postes


def controles_dpgf(c: Consultation) -> list[Constat]:
    out = []
    k = c.coefficients
    manquants = [n for n, v in (("taux horaire", k.taux_horaire), ("frais généraux", k.frais_generaux), ("aléas", k.aleas), ("marge", k.marge)) if v is None]
    if manquants:
        out.append(Constat(gravite="a_verifier", objet="Coefficients", message="À renseigner par SIE : " + ", ".join(manquants)))
    for p in c.dpgf:
        for a in p.alertes:
            out.append(Constat(gravite="a_verifier", objet=f"Poste {p.code or ''} {p.designation}".strip(), message=a))
    devis_non_associes = sum(1 for d in c.devis for l in d.lignes if not l.ligne_quantite_id)
    if devis_non_associes:
        out.append(Constat(gravite="info", objet="Devis", message=f"{devis_non_associes} lignes de devis sans poste associé"))
    return out


def exporter_dpgf(c: Consultation, chemin: str | Path, modele: str | Path | None = None) -> Path:
    """DPGF Excel avec formules. Avec un modèle SIE, remplit ses colonnes Quantité / PU en reconnaissant les désignations."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    chemin = Path(chemin)
    jaune = PatternFill("solid", fgColor="FFF2CC")
    if modele:
        wb = load_workbook(modele)
        restants = list(c.dpgf)
        for ws in wb.worksheets:
            rangees = [[cell.value for cell in row] + [None] * 5 for row in ws.iter_rows()]
            for i, r in enumerate(rangees[:30]):
                cols = _colonnes(r)
                prix = next((k for k, v in enumerate(r) if any(re.search(m, normaliser(str(v or ""))) for m in ENTETES_PRIX)), None)
                if "designation" not in cols or prix is None:
                    continue
                for j in range(i + 1, ws.max_row):
                    des = ws.cell(row=j + 1, column=cols["designation"] + 1).value
                    if not des or not restants:
                        continue
                    best = max(restants, key=lambda p: similarite(str(des), p.designation))
                    if similarite(str(des), best.designation) < 0.5:
                        continue
                    if "quantite" in cols and best.quantite is not None:
                        ws.cell(row=j + 1, column=cols["quantite"] + 1, value=best.quantite)
                    cell = ws.cell(row=j + 1, column=prix + 1, value=best.prix_unitaire_vente)
                    if best.prix_unitaire_vente is None or best.alertes:
                        cell.fill = jaune
                    restants.remove(best)
                break
        wb.save(chemin)
        if restants:
            c.journal.append(f"F5 : {len(restants)} postes non retrouvés dans le modèle de DPGF SIE.")
        return chemin

    wb = Workbook()
    ws = wb.active
    ws.title = "DPGF"
    k = c.coefficients
    ws.append(["Coefficients SIE (à saisir ou valider)"])
    ws["A1"].font = Font(bold=True)
    params = [("Taux horaire (€/h)", k.taux_horaire), ("Frais généraux (%)", k.frais_generaux), ("Aléas (%)", k.aleas), ("Marge (%)", k.marge)]
    for nom, v in params:
        ws.append([nom, v])
        if v is None:
            ws.cell(row=ws.max_row, column=2).fill = jaune
    ws.append([])
    entete = ["Lot", "N°", "Désignation", "Unité", "Quantité", "Prix achat HT", "Fournisseur", "Temps pose (h/u)",
              "Déboursé unitaire", "PU vente HT", "Total HT", "Alertes"]
    ws.append(entete)
    h = ws.max_row
    for cell in ws[h]:
        cell.font = Font(bold=True)
    for p in c.dpgf:
        r = ws.max_row + 1
        ws.append([p.lot, p.code, p.designation, p.unite, p.quantite, p.prix_materiel, p.fournisseur, p.temps_pose_h,
                   f"=IF(F{r}=\"\",\"\",F{r}+N(H{r})*$B$2)",
                   f"=IF(OR(I{r}=\"\",$B$3=\"\",$B$4=\"\",$B$5=\"\"),\"\",ROUND(I{r}*(1+$B$3/100)*(1+$B$4/100)*(1+$B$5/100),2))",
                   f"=IF(OR(J{r}=\"\",E{r}=\"\"),\"\",ROUND(J{r}*E{r},2))",
                   " ; ".join(p.alertes)])
        if p.alertes:
            for cell in ws[r]:
                cell.fill = jaune
    fin = ws.max_row
    ws.append([None, None, "TOTAL HT", None, None, None, None, None, None, None, f"=SUM(K{h + 1}:K{fin})"])
    ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
    for col, w in zip("ABCDEFGHIJKL", (6, 8, 45, 7, 9, 12, 16, 10, 12, 12, 12, 50)):
        ws.column_dimensions[col].width = w
    wb.save(chemin)
    return chemin

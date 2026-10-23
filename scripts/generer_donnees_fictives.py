"""Génère un DCE et des archives FICTIFS pour tester l'outil sans donnée SIE.

    python scripts/generer_donnees_fictives.py exemples/

Tous les noms, adresses, montants et dates sont inventés.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import docx
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

STYLES = getSampleStyleSheet()


def pdf(chemin: Path, blocs: list[str]) -> None:
    """Chaque bloc commençant par '# ' est un titre ; '---' force un saut de page."""
    doc = SimpleDocTemplate(str(chemin), pagesize=A4)
    elements = []
    for b in blocs:
        if b == "---":
            elements.append(PageBreak())
        elif b.startswith("# "):
            elements.append(Paragraph(b[2:], STYLES["Heading2"]))
        else:
            for ligne in b.split("\n"):
                elements.append(Paragraph(ligne, STYLES["BodyText"]))
            elements.append(Spacer(1, 6))
    doc.build(elements)


RC = [
    "# REGLEMENT DE LA CONSULTATION (RC)",
    "Commune fictive de Valbourg - Réhabilitation du groupe scolaire des Tilleuls\nRéférence de la consultation : VAL-2026-014",
    "# Article 1 - Objet de la consultation",
    "Objet de la consultation : Réhabilitation des installations techniques du groupe scolaire des Tilleuls\n"
    "Lieu d'exécution : 12 rue des Tilleuls, 44999 Valbourg\n"
    "Maître d'ouvrage : Commune de Valbourg, Hôtel de ville, 1 place de la Mairie, 44999 Valbourg\n"
    "Maître d'oeuvre : Cabinet fictif Atelier Durand Ingénierie, 8 quai Ouest, 44000 Nantes\n"
    "Contrôleur technique : Bureau fictif Contrôle Ouest\n"
    "Coordonnateur SPS : Cabinet fictif Prévention Chantiers 44",
    "# Article 2 - Procédure et décomposition en lots",
    "La présente consultation est passée selon une procédure adaptée ouverte. Les travaux sont répartis en lots :",
    "Lot 1 : Gros oeuvre - démolitions\nLot 2 : Electricité courants forts et courants faibles\nLot 3 : Plomberie sanitaires\nLot 4 : Chauffage ventilation climatisation",
    "# Article 3 - Visite des lieux",
    "La visite des lieux est obligatoire sous peine d'irrecevabilité de l'offre. Elle aura lieu le mardi 10 novembre 2026 à 10h00. "
    "Une attestation de visite sera remise aux candidats et devra être jointe à l'offre.",
    "---",
    "# Article 4 - Contenu des plis",
    "Les candidats doivent fournir les pièces suivantes :",
    "- la lettre de candidature DC1 et la déclaration du candidat DC2 ;\n"
    "- les attestations d'assurance responsabilité civile et décennale en cours de validité ;\n"
    "- l'attestation de vigilance URSSAF de moins de six mois ;\n"
    "- les certificats de qualification Qualifelec ou équivalent ;\n"
    "- une liste de références de travaux similaires sur les cinq dernières années.",
    "L'offre comprendra l'acte d'engagement daté et signé, la DPGF complétée, le mémoire technique et le planning prévisionnel.",
    "# Article 5 - Jugement des offres",
    "Les offres seront jugées selon les critères pondérés suivants :",
    "1. Prix des prestations : 40 %\n2. Valeur technique : 60 %\n"
    "- Méthodologie d'intervention en site occupé : 25 %\n- Moyens humains et matériels affectés au chantier : 15 %\n"
    "- Gestion environnementale du chantier et des déchets : 10 %\n- Planning et phasage des travaux : 10 %",
    "# Article 6 - Conditions de remise des offres",
    "Date limite de réception des offres : vendredi 27 novembre 2026 à 12h00.\n"
    "Les questions des candidats devront parvenir au plus tard le 18 novembre 2026 via le profil d'acheteur.\n"
    "Le délai de validité des offres est fixé à 120 jours à compter de la date limite de réception.\n"
    "Contact administratif : marches@valbourg-fictif.fr - Contact technique : atelier.durand@exemple-fictif.fr",
]

CCAP = [
    "# CAHIER DES CLAUSES ADMINISTRATIVES PARTICULIERES (CCAP)",
    "Commune fictive de Valbourg - Groupe scolaire des Tilleuls",
    "# Article 7 - Délais d'exécution",
    "Le délai d'exécution global est de 5 mois, période de préparation de 1 mois comprise. Les travaux bruyants sont réalisés pendant les vacances scolaires.",
    "# Article 8 - Clauses sociales et environnementales",
    "8.1 Clause d'insertion. Le titulaire du lot 2 réalise au minimum 70 heures d'insertion professionnelle au bénéfice de personnes éloignées de l'emploi.\n"
    "8.2 Gestion des déchets. Le titulaire assure le tri des déchets de chantier sur site et fournit les bordereaux de suivi des déchets dangereux.\n"
    "8.3 Nuisances. Le titulaire limite le bruit et les poussières, l'établissement restant en activité pendant les travaux.",
    "# Article 9 - Pénalités",
    "Les pièces non remises dans les délais seront rejetées. Le dossier des ouvrages exécutés est remis sous peine de retenue.",
]

CCTP_LOT2 = [
    "# CCTP - LOT 2 : ELECTRICITE COURANTS FORTS ET COURANTS FAIBLES",
    "# Article 1 - Généralités",
    "Les installations sont réalisées conformément à la norme NF C 15-100 et au règlement de sécurité ERP. Les travaux se déroulent en site occupé, par phases successives validées avec la direction de l'école.",
    "# Article 2 - Limites de prestations",
    "L'alimentation électrique des centrales de traitement d'air est à la charge du lot 2 jusqu'aux bornes du coffret fourni par le lot 4. "
    "Les réservations dans les cloisons sont à la charge du lot 1, sur plans fournis par le lot 2.",
    "# Article 3 - Sécurité incendie",
    "Le lot 2 fournit et pose l'équipement d'alarme de type 4 et assure les asservissements du désenfumage en coordination avec le lot 4.",
    "# Article 4 - Eclairage",
    "Remplacement de l'ensemble des luminaires par des luminaires LED avec détection de présence, afin de réduire les consommations d'énergie.",
]

CCTP_LOT4 = [
    "# CCTP - LOT 4 : CHAUFFAGE VENTILATION CLIMATISATION",
    "# Article 1 - Objet",
    "Remplacement des centrales de traitement d'air et des émetteurs de chauffage.",
    "# Article 2 - Interfaces",
    "Le coffret de commande des centrales est fourni par le lot 4 ; son alimentation est à la charge du lot 2. Les clapets coupe-feu sont asservis au SSI en coordination avec le lot 2.",
]

ARCHIVES = {
    "2024_ecole_maternelle_fictive_Orvault": {
        "meta": {"activite": "CFO/CFA", "annee": 2024, "type": "enseignement", "resultat": "gagné"},
        "memoire": [
            "# MEMOIRE TECHNIQUE - Ecole maternelle (affaire fictive)",
            "# 1. Méthodologie d'intervention en site occupé",
            "Nos équipes interviennent en site occupé en respectant les horaires de l'école. Les travaux bruyants sont planifiés le mercredi et pendant les vacances scolaires. "
            "Chaque zone de travaux est isolée par des barrières et un balisage, et un point hebdomadaire est tenu avec la direction de l'établissement.",
            "# 2. Gestion environnementale du chantier",
            "Les déchets sont triés sur site dans des bennes séparées (DIB, métaux, DEEE). Les luminaires déposés sont repris par la filière agréée et les bordereaux de suivi sont transmis au maître d'ouvrage.",
            "# 3. Moyens humains",
            "Un chargé d'affaires, un chef de chantier habilité B2V BR et deux électriciens qualifiés sont affectés au chantier.",
        ],
        "attestations": {"attestation_qualifelec.pdf": "Attestation de qualification Qualifelec. Valable jusqu'au 31/12/2027."},
    },
    "2023_ehpad_fictif_Reze": {
        "meta": {"activite": "CFO/SSI", "annee": 2023, "type": "santé", "resultat": "gagné"},
        "memoire": [
            "# MEMOIRE TECHNIQUE - EHPAD (affaire fictive)",
            "# 1. Organisation en milieu occupé",
            "L'établissement restant ouvert, nous travaillons chambre par chambre avec un phasage validé chaque semaine. Les coupures électriques sont annoncées 48 heures à l'avance.",
            "# 2. Sécurité incendie",
            "Le remplacement du SSI est réalisé sans interruption de la mise en sécurité : l'ancien et le nouveau système fonctionnent en parallèle jusqu'à la réception.",
        ],
        "attestations": {"attestation_assurance_decennale.pdf": "Attestation d'assurance décennale. Date de fin de validité : 31/03/2026."},
    },
    "2022_bureaux_fictifs_Nantes": {
        "meta": {"activite": "CFA", "annee": 2022, "type": "tertiaire", "resultat": "perdu"},
        "memoire": [
            "# MEMOIRE TECHNIQUE - Bureaux (affaire fictive)",
            "# 1. Câblage courants faibles",
            "Le câblage VDI est réalisé en catégorie 6A avec recette de chaque prise.",
        ],
        "attestations": {},
    },
}


def generer(racine: Path) -> None:
    dce = racine / "dce_fictif_valbourg"
    arch = racine / "archives_fictives"
    dce.mkdir(parents=True, exist_ok=True)
    pdf(dce / "01_RC_Valbourg.pdf", RC)
    pdf(dce / "02_CCAP.pdf", CCAP)
    pdf(dce / "03_CCTP_lot02_electricite.pdf", CCTP_LOT2)
    pdf(dce / "04_CCTP_lot04_CVC.pdf", CCTP_LOT4)
    d = docx.Document()
    d.add_heading("ACTE D'ENGAGEMENT", 0)
    d.add_paragraph("Marché n° VAL-2026-014. Le candidat s'engage sur le montant de son offre. Document daté et signé.")
    d.save(dce / "05_AE.docx")
    wb = Workbook()
    ws = wb.active
    ws.title = "DPGF lot 2"
    ws.append(["N°", "Désignation", "Unité", "Quantité", "PU HT", "Total HT"])
    ws.append(["2.1", "Luminaire LED 600x600", "u", None, None, None])
    ws.append(["2.2", "Câble U1000 R2V 3G2,5", "ml", None, None, None])
    wb.save(dce / "06_DPGF_lot02.xlsx")
    (dce / "07_plan_RDC_lot02.dwg").write_bytes(b"fichier DWG fictif")

    for nom, a in ARCHIVES.items():
        rep = arch / nom
        rep.mkdir(parents=True, exist_ok=True)
        (rep / "affaire.json").write_text(json.dumps(a["meta"], ensure_ascii=False), encoding="utf-8")
        pdf(rep / "memoire_technique.pdf", a["memoire"])
        for fichier, texte in a["attestations"].items():
            pdf(rep / fichier, [texte])
    # Données du groupe 2 : métré du bureau d'études, devis fournisseurs, pièces de réponse
    g2 = racine / "chiffrage_fictif"
    g2.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lot 2 electricite"
    ws.append(["Métré lot 2 - Groupe scolaire des Tilleuls (fictif)"])
    ws.append([])
    ws.append(["N°", "Désignation", "U", "Qté"])
    ws.append([None, "ECLAIRAGE", None, None])
    ws.append(["2.1", "Luminaire LED 600x600 encastré 36W", "u", 84])
    ws.append(["2.2", "Détecteur de présence plafonnier", "u", 22])
    ws.append(["2.3", "Bloc autonome d'éclairage de sécurité BAES", "u", 18])
    ws.append([None, "DISTRIBUTION", None, None])
    ws.append(["2.4", "Câble U1000 R2V 3G2,5", "ml", 1450])
    ws.append(["2.5", "Chemin de câble dalle marine 100 mm", "ml", 210])
    ws.append(["2.6", "Tableau divisionnaire étage", "u", 2])
    ws.append([None, "SECURITE INCENDIE", None, None])
    ws.append(["2.7", "Déclencheur manuel alarme type 4", "u", 12])
    ws.append(["2.8", "Sirène d'alarme incendie", "u", "à compter"])
    ws.append(["2.9", "Dépose des luminaires existants", "ft", 1])
    wb.save(g2 / "metre_lot02.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.append(["Devis n° 4521 - Elec Distribution Ouest (fictif)"])
    ws.append(["Référence", "Désignation", "Unité", "Quantité", "Prix unitaire HT"])
    ws.append(["LED600", "Dalle LED 600x600 36W encastrée", "u", 84, 42.50])
    ws.append(["DET-P", "Détecteur présence plafond 360°", "u", 22, 38.00])
    ws.append(["BAES", "BAES évacuation 45 lm", "u", 18, 29.90])
    ws.append(["R2V325", "Câble R2V 3G2,5 U1000", "m", 1000, 1.12])
    ws.append(["CDC100", "Chemin de câbles dalle marine 100", "ml", 210, 9.80])
    wb.save(g2 / "devis_elec_distribution_ouest.xlsx")

    pdf(g2 / "devis_lumiere_plus.pdf", [
        "# DEVIS 2026-118 - Lumière Plus (fictif)",
        "Dalle LED 600x600 36W 84 u 39,90 3 351,60",
        "Detecteur de presence plafonnier 22 u 41,00 902,00",
        "Bloc BAES evacuation 18 u 31,50 567,00",
    ])

    rep = g2 / "pieces_reponse"
    rep.mkdir(exist_ok=True)
    for nom in ("DC1_lettre_candidature.pdf", "DC2_declaration_candidat.pdf", "attestation_URSSAF_2026.pdf",
                "certificat_Qualifelec.pdf", "references_travaux_similaires.pdf", "acte_engagement_signe.pdf",
                "planning_previsionnel.pdf"):
        pdf(rep / nom, [f"Document fictif : {nom}"])
    print(f"DCE fictif : {dce}\nArchives fictives : {arch}\nChiffrage fictif : {g2}")


if __name__ == "__main__":
    generer(Path(sys.argv[1] if len(sys.argv) > 1 else "exemples"))

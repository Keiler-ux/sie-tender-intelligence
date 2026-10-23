"""Points du cahier des charges : interfaces inter-lots, synthèse des lots retenus, cadre de réponse, modèle de fiche."""
from __future__ import annotations

import sys
from pathlib import Path

import docx
import pytest
from openpyxl import Workbook, load_workbook

RACINE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RACINE))
sys.path.insert(0, str(RACINE / "scripts"))

from sie_tender.export import complements_a_demander, exporter_fiche_modele, exporter_trame_docx  # noqa: E402
from sie_tender.f2_synthese import exporter_synthese_docx, heures_insertion, matrice_interfaces, synthese_lot  # noqa: E402
from sie_tender.f6_archives import Bibliotheque  # noqa: E402
from sie_tender.ingestion import importer_dce  # noqa: E402
from sie_tender.pipeline import analyser_dce  # noqa: E402


@pytest.fixture(scope="module")
def racine(tmp_path_factory):
    from generer_dce_test import generer as generer_test
    from generer_donnees_fictives import generer
    d = tmp_path_factory.mktemp("cdc")
    generer(d)
    generer_test(d)
    return d


def test_activites_et_interfaces(racine):
    c = analyser_dce(racine / "dce_fictif_valbourg")
    assert {l.numero: l.activites for l in c.lots}["4"] == ["CVC"]
    assert "CFO" in c.lots[1].activites
    c.lots[1].retenu = True  # lot 2 seulement
    inter = matrice_interfaces(c)
    # une interface écrite dans le CCTP du lot 4 (non retenu) touche le lot 2 retenu
    assert any(i.origine == "4" and "2" in i.concernes and i.impacte_retenu for i in inter)
    assert [h for h, _ in heures_insertion(c.exigences)] == [70]


def test_synthese_lot_retenu(racine, tmp_path):
    dce = racine / "dce_hopital_3_lots"
    c = analyser_dce(dce)
    c.fiche.decision = "GO"
    c.lots[1].retenu = True
    pieces = importer_dce(dce)
    s = synthese_lot(c, c.lots[1], pieces)
    assert s.prestations and s.prestations[0].source.fichier == "02_CCTP_Lot_02.pdf"
    assert any("pénalité" in e.texte for e in s.risques)
    assert any("collaborateur chargé de la visite" in q for q in s.questions)
    exporter_synthese_docx(c, pieces, tmp_path / "s.docx")
    texte = "\n".join(p.text for p in docx.Document(tmp_path / "s.docx").paragraphs)
    assert "Lot 2" in texte and "Lot 1 :" not in texte  # seuls les lots retenus


def test_cadre_de_reponse_prioritaire(racine, tmp_path):
    b = Bibliotheque.indexer(racine / "archives_fictives")
    c = analyser_dce(racine / "dce_hopital_3_lots", b)
    titres = [s.titre for s in c.trame_memoire]
    assert titres[0].startswith("Chapitre 1") and "1.1 Phasage par groupes de chambres" in titres
    assert all(s.impose_par_acheteur for s in c.trame_memoire)
    assert not any("pages" in t for t in titres)
    assert c.references and c.references[0].type == "santé"
    s = c.trame_memoire[0]
    s.contenu, s.sources_reprises = "Texte repris. [À COMPLÉTER : nombre de chambres par phase]", [s.passages_proposes[0].source]
    assert ("Chapitre 1 - Organisation du chantier en établissement de santé occupé", "nombre de chambres par phase") in complements_a_demander(c)
    exporter_trame_docx(c, tmp_path / "m.docx")
    texte = "\n".join(p.text for p in docx.Document(tmp_path / "m.docx").paragraphs)
    assert "Sources des contenus repris" in texte and "Références d'affaires comparables" in texte


def test_sans_cadre_plan_des_criteres(racine):
    c = analyser_dce(racine / "dce_fictif_valbourg")
    assert c.cadre_reponse is None
    assert c.trame_memoire[0].titre == "Méthodologie d'intervention en site occupé"


def test_fiche_au_modele_sie(racine, tmp_path):
    c = analyser_dce(racine / "dce_hopital_3_lots")
    c.fiche.decision, c.fiche.responsable_visite = "GO", "Camille Martin"
    c.lots[1].retenu = True
    wb = Workbook()
    ws = wb.active
    ws.append(["Référence", "{{reference}}"])
    ws.append(["{{lot_numero}}", "{{lot_intitule}}", "{{lot_retenu}}"])
    ws.append(["Visite", "{{responsable_visite}}", "{{balise_inconnue}}"])
    wb.save(tmp_path / "modele.xlsx")
    assert exporter_fiche_modele(c, tmp_path / "modele.xlsx", tmp_path / "fiche.xlsx") == ["balise_inconnue"]
    lignes = list(load_workbook(tmp_path / "fiche.xlsx").active.iter_rows(values_only=True))
    assert lignes[0][1] == "CHSA-TX-2026-31"
    assert [r[2] for r in lignes[1:4]] == ["☐", "☒", "☐"]
    assert lignes[4][1] == "Camille Martin"


def test_evaluation():
    import evaluer
    assert evaluer.main([]) == 0

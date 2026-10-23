"""Exports : JSON (format d'échange avec le groupe 2), Excel, Markdown et Word."""
from __future__ import annotations

import json
import re
from pathlib import Path

from .models import Consultation, Source, Statut

LIBELLES = {
    "reference": "Référence", "intitule": "Intitulé", "objet": "Objet", "lieu": "Lieu d'exécution",
    "maitre_ouvrage": "Maître d'ouvrage", "maitre_oeuvre": "Maître d'œuvre", "procedure": "Procédure",
    "date_limite_remise": "Date et heure limites de remise", "date_limite_questions": "Date limite des questions",
    "visite": "Visite", "visite_obligatoire": "Visite obligatoire", "contacts": "Contacts",
    "duree_validite_offres": "Délai de validité des offres",
    "intervenants": "Autres intervenants (BET, contrôle, SPS, OPC)",
}
STATUTS = {Statut.EXTRAIT: "extrait", Statut.A_VERIFIER: "À VÉRIFIER", Statut.CORRIGE: "corrigé par SIE", Statut.VALIDE: "validé"}


def ref(src: Source | None) -> str:
    if src is None:
        return ""
    morceaux = [src.fichier]
    if src.page:
        morceaux.append(f"p. {src.page}")
    if src.article:
        morceaux.append(f"art. {src.article}")
    return ", ".join(morceaux)


def exporter_json(c: Consultation, chemin: Path) -> None:
    chemin.write_text(c.model_dump_json(indent=2), encoding="utf-8")


def exporter_excel(c: Consultation, chemin: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    jaune = PatternFill("solid", fgColor="FFF2CC")

    def feuille(titre, entetes, lignes, largeurs):
        ws = wb.create_sheet(titre)
        ws.append(entetes)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for l in lignes:
            ws.append(l)
            if any(isinstance(v, str) and v.startswith("À VÉRIFIER") for v in l):
                for cell in ws[ws.max_row]:
                    cell.fill = jaune
        for i, w in enumerate(largeurs):
            ws.column_dimensions[chr(65 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"

    wb.remove(wb.active)
    f = c.fiche
    lignes = [[LIBELLES.get(k, k), ch.valeur or "", STATUTS[ch.statut], ref(ch.source), ch.source.extrait if ch.source else ""]
              for k, ch in f.champs().items()]
    lignes.append(["Lots", "\n".join(f"Lot {l.numero} : {l.intitule}" for l in c.lots), "extrait" if c.lots else "À VÉRIFIER", "", ""])
    lignes.append(["Alertes", "\n".join(f.alertes), "", "", ""])
    lignes.append(["Décision GO / NO GO", f.decision or "(à saisir par la direction)",
                   f"{f.decideur} le {f.date_decision:%d/%m/%Y}" if f.date_decision else "", "", ""])
    lignes.append(["Lots retenus", ", ".join(f"lot {l.numero}" for l in c.lots if l.retenu) or "(après GO, sélection par SIE)", "", "", ""])
    lignes.append(["Responsable de la visite", f.responsable_visite or "(à saisir par SIE)", "", "", ""])
    lignes.append(["Date de la visite", f.date_visite or "(à saisir par SIE)", "", "", ""])
    feuille("Fiche", ["Champ", "Valeur", "Statut", "Source", "Extrait"], lignes, [28, 50, 14, 30, 70])
    feuille("Lots", ["N°", "Intitulé", "Activités SIE", "Retenu", "Responsable", "Source"],
            [[l.numero, l.intitule, ", ".join(l.activites), {True: "oui", False: "non", None: "(décision SIE)"}[l.retenu],
              l.responsable or "", ref(l.source)] for l in c.lots], [6, 50, 18, 16, 20, 40])
    from .f2_synthese import matrice_interfaces, tableau_croise
    inter = matrice_interfaces(c)
    axes, grille = tableau_croise(inter, c.lots)
    feuille("Matrice interfaces", ["Lot d'origine \\ lot concerné"] + [f"Lot {l.numero}" for l in c.lots],
            [[o if o == "commun" else f"Lot {o}"] + [grille[o][l.numero] or "" for l in c.lots] for o in axes],
            [26] + [10] * len(c.lots))
    feuille("Interfaces", ["Origine", "Lots concernés", "Nature", "Impact lot retenu", "Texte du DCE", "Source"],
            [[i.origine, ", ".join(i.concernes), ", ".join(i.themes), "oui" if i.impacte_retenu else "", i.texte, ref(i.source)]
             for i in inter], [10, 14, 26, 12, 80, 34])
    feuille("Exigences", ["Lot", "Catégorie", "Précision", "Exigence", "Statut", "Source"],
            [[e.lot or "commun", e.categorie.value, e.sous_categorie or "", e.texte, STATUTS[e.statut], ref(e.source)] for e in c.exigences],
            [8, 14, 22, 90, 12, 34])
    crit = []
    for cr in c.criteres:
        crit.append([cr.intitule, cr.ponderation, "", STATUTS[cr.statut], ref(cr.source)])
        for sc in cr.sous_criteres:
            crit.append([f"   {sc.intitule}", "", sc.ponderation, STATUTS[sc.statut], ref(sc.source)])
    feuille("Critères", ["Critère", "Pondération (%)", "Sous-pondération", "Statut", "Source"], crit, [50, 16, 16, 12, 34])
    feuille("Pièces à fournir", ["Étape", "Pièce", "Document proposé (archives)", "Validité", "Statut", "Source"],
            [[p.etape, p.intitule, p.document_propose or "", p.validite.strftime("%d/%m/%Y") if p.validite else "", STATUTS[p.statut], ref(p.source)]
             for p in c.pieces_a_fournir], [12, 40, 40, 12, 12, 34])
    feuille("Pièces du DCE", ["Fichier", "Type", "Lot", "Pages", "Lue", "Remarque"],
            [[p.fichier, p.type.value, p.lot or "commune", p.nb_pages, "oui" if p.lisible else "non", p.remarque or ""] for p in c.pieces],
            [40, 10, 10, 8, 8, 60])
    feuille("Journal", ["Hypothèses, erreurs et limites observées"], [[j] for j in c.journal], [120])
    wb.save(chemin)


def exporter_markdown(c: Consultation, chemin: Path) -> None:
    f = c.fiche
    L = [f"# Analyse du DCE : {f.intitule.valeur or f.objet.valeur or c.dossier}", ""]
    L += ["## Fiche de renseignement (F1)", "", "| Champ | Valeur | Statut | Source |", "| --- | --- | --- | --- |"]
    for k, ch in f.champs().items():
        L.append(f"| {LIBELLES.get(k, k)} | {(ch.valeur or '').replace('|', '/')} | {STATUTS[ch.statut]} | {ref(ch.source)} |")
    L += ["", "**Lots** : " + ("; ".join(f"lot {l.numero} {l.intitule}" for l in c.lots) or "à vérifier"), ""]
    if f.alertes:
        L += ["**Alertes**", ""] + [f"- {a}" for a in f.alertes] + [""]
    L += ["**Décision GO / NO GO** : à saisir par la direction SIE.", ""]
    L += ["## Critères de notation (F2)", ""]
    for cr in c.criteres:
        L.append(f"- {cr.intitule} : {cr.ponderation} % ({ref(cr.source)})")
        L += [f"    - {sc.intitule} : {sc.ponderation} %" for sc in cr.sous_criteres]
    L += ["", "## Exigences par catégorie (F2)", ""]
    for cat in sorted({e.categorie.value for e in c.exigences}):
        L += [f"### {cat}", ""]
        for e in [e for e in c.exigences if e.categorie.value == cat]:
            L.append(f"- [lot {e.lot or 'commun'}] {e.texte} *({ref(e.source)})*")
        L.append("")
    L += ["## Pièces à fournir", ""]
    for p in c.pieces_a_fournir:
        doc = f" → {p.document_propose}" if p.document_propose else ""
        L.append(f"- [{p.etape}] {p.intitule}{doc} ({STATUTS[p.statut]}, {ref(p.source)})")
    L += ["", "## Trame du mémoire technique (F6)", ""]
    for s in c.trame_memoire:
        pond = f" ({s.ponderation} %)" if s.ponderation else ""
        L.append(f"### {s.titre}{pond} : {s.statut.value}")
        if s.note:
            L.append(f"_{s.note}_")
        for pa in s.passages_proposes[:3]:
            L.append(f"- {pa.texte[:300]}… *({ref(pa.source)}, score {pa.score})*")
        L.append("")
    L += ["## Journal", ""] + [f"- {j}" for j in c.journal]
    chemin.write_text("\n".join(L), encoding="utf-8")


def exporter_trame_docx(c: Consultation, chemin: Path) -> None:
    import docx
    from docx.shared import Pt, RGBColor

    d = docx.Document()

    def note(texte: str, couleur: RGBColor | None = None, gras: bool = False) -> None:
        r = d.add_paragraph().add_run(texte)
        r.font.size, r.italic, r.bold = Pt(8), not gras, gras
        if couleur:
            r.font.color.rgb = couleur

    orange = RGBColor(0xC0, 0x50, 0x00)
    d.add_heading(f"Mémoire technique : {c.fiche.objet.valeur or c.fiche.intitule.valeur or 'consultation'}", 0)
    plan = (f"Plan repris du cadre de réponse de l'acheteur ({ref(c.cadre_reponse)})." if c.cadre_reponse
            else "Plan construit à partir des critères de notation du RC.")
    p = d.add_paragraph("Projet généré par l'outil d'aide à la réponse. " + plan + " Chaque passage est une proposition à relire, "
                        "adapter et valider ; les sources sont indiquées entre crochets.")
    p.runs[0].italic = True
    for s in c.trame_memoire:
        titre = s.titre + (f" ({s.ponderation:g} %)" if s.ponderation else "")
        d.add_heading(titre, 2 if re.match(r"\d+\.\d", s.titre) else 1)
        if s.contenu:
            for para in s.contenu.split("\n\n"):
                d.add_paragraph(para.strip())
            if s.sources_reprises:
                note("[Sources des contenus repris : " + " ; ".join(ref(x) for x in s.sources_reprises) + "]")
            continue
        if s.note or s.statut.value == "a_completer":
            note(f"[À COMPLÉTER] {s.note or ''}".strip(), orange, gras=True)
        for pa in s.passages_proposes[:3]:
            d.add_paragraph(pa.texte)
            note(f"[Source : {ref(pa.source)}]")
    refs = [r for r in c.references if r.retenue] or [r for r in c.references if r.retenue is None]
    if refs:
        d.add_heading("Références d'affaires comparables", 1)
        if not any(r.retenue for r in c.references):
            note("[À VALIDER] Affaires proposées par l'outil d'après les archives : SIE choisit celles à citer.", orange, gras=True)
        for r in refs:
            details = ", ".join(str(x) for x in (r.annee, r.type, r.activite) if x)
            d.add_paragraph(f"{r.affaire}" + (f" ({details})" if details else ""), style="List Bullet")
    if c.pieces_jointes_memoire:
        d.add_heading("Pièces jointes au mémoire", 1)
        for f in c.pieces_jointes_memoire:
            d.add_paragraph(f, style="List Bullet")
    d.save(chemin)


_A_COMPLETER = re.compile(r"\[À COMPLÉTER\s*:?\s*([^\]]*)\]")


def complements_a_demander(c: Consultation) -> list[tuple[str, str]]:
    """(section, ce qui manque) : ce que le bureau d'études ou le chargé d'affaires doit fournir."""
    out: list[tuple[str, str]] = []
    for s in c.trame_memoire:
        if s.contenu:
            out += [(s.titre, m.strip() or "à compléter") for m in _A_COMPLETER.findall(s.contenu)]
        elif s.statut.value == "a_completer":
            out.append((s.titre, s.note or "section à rédiger"))
        elif s.statut.value != "valide":
            out.append((s.titre, "passages proposés à relire, adapter et valider"))
    if c.references and not any(r.retenue for r in c.references):
        out.append(("Références", "choisir les affaires comparables à citer"))
    exige = re.compile(r"planning|note de calcul|synoptique")
    for pf in c.pieces_a_fournir:
        if exige.search(pf.intitule.lower()) and not any(exige.search(f.lower()) for f in c.pieces_jointes_memoire):
            out.append(("Pièces jointes", f"{pf.intitule} (exigé : {ref(pf.source)})"))
    return out


def exporter_complements_docx(c: Consultation, chemin: Path) -> None:
    import docx

    d = docx.Document()
    d.add_heading("Compléments à fournir pour le mémoire technique", 0)
    d.add_paragraph(f"Affaire : {c.fiche.intitule.valeur or c.fiche.objet.valeur or c.id}")
    t = d.add_table(rows=1, cols=3)
    t.style = "Light Grid Accent 1"
    for cell, txt in zip(t.rows[0].cells, ("Section", "À fournir", "Fait")):
        cell.text = txt
    for section, quoi in complements_a_demander(c) or [("—", "Rien à compléter")]:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = section, quoi, "☐"
    d.save(chemin)


def exporter_tout(c: Consultation, dossier: str | Path) -> dict[str, Path]:
    dossier = Path(dossier)
    dossier.mkdir(parents=True, exist_ok=True)
    chemins = {
        "json": dossier / "consultation.json",
        "excel": dossier / "analyse_dce.xlsx",
        "markdown": dossier / "rapport.md",
        "memoire": dossier / "trame_memoire.docx",
    }
    exporter_json(c, chemins["json"])
    exporter_excel(c, chemins["excel"])
    exporter_markdown(c, chemins["markdown"])
    exporter_trame_docx(c, chemins["memoire"])
    return chemins


def valeurs_fiche(c: Consultation) -> dict[str, str]:
    """Valeurs disponibles pour remplir le modèle de fiche SIE : {{reference}}, {{lots}}, {{decision}}..."""
    f = c.fiche
    v = {k: (ch.valeur or "À VÉRIFIER") for k, ch in f.champs().items()}
    v.update({
        "lots": "\n".join(f"Lot {l.numero} : {l.intitule}" for l in c.lots) or "À VÉRIFIER",
        "lots_retenus": ", ".join(f"Lot {l.numero}" for l in c.lots if l.retenu),
        "decision": f.decision or "", "decideur": f.decideur or "",
        "date_decision": f"{f.date_decision:%d/%m/%Y}" if f.date_decision else "",
        "responsable_visite": f.responsable_visite or "", "date_visite": f.date_visite or "",
        "alertes": "\n".join(f.alertes),
    })
    return v


_BALISE = re.compile(r"\{\{\s*(\w+)\s*\}\}")


def exporter_fiche_modele(c: Consultation, modele: Path, chemin: Path) -> list[str]:
    """Remplit le modèle Excel de fiche fourni par SIE (balises {{champ}}) ; renvoie les balises inconnues.

    Les lignes de lots peuvent aussi se dupliquer : une ligne contenant {{lot_numero}} / {{lot_intitule}}
    / {{lot_retenu}} est répétée pour chaque lot, avec une case « ☒ / ☐ ».
    """
    from copy import copy

    from openpyxl import load_workbook

    wb = load_workbook(modele)
    valeurs = valeurs_fiche(c)
    inconnues: set[str] = set()
    for ws in wb.worksheets:
        ligne_lot = next((row[0].row for row in ws.iter_rows() if any(isinstance(x.value, str) and "{{lot_" in x.value for x in row)), None)
        if ligne_lot and c.lots:
            ws.insert_rows(ligne_lot + 1, len(c.lots) - 1)
            modele_cells = [(x.column, x.value, copy(x._style)) for x in ws[ligne_lot]]
            for i, lot in enumerate(c.lots):
                vl = {"lot_numero": lot.numero, "lot_intitule": lot.intitule, "lot_retenu": "☒" if lot.retenu else "☐",
                      "lot_responsable": lot.responsable or "", "lot_activites": ", ".join(lot.activites)}
                for col, val, style in modele_cells:
                    dst = ws.cell(row=ligne_lot + i, column=col)
                    dst._style = copy(style)
                    dst.value = _BALISE.sub(lambda m: vl.get(m.group(1), m.group(0)), val) if isinstance(val, str) else val
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    inconnues |= {m for m in _BALISE.findall(cell.value) if m not in valeurs}
                    cell.value = _BALISE.sub(lambda m: valeurs.get(m.group(1), m.group(0)), cell.value)
    wb.save(chemin)
    return sorted(inconnues)

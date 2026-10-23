"""Génère trois DCE FICTIFS supplémentaires pour tester l'outil, chacun en dossier et en ZIP.

    python scripts/generer_dce_test.py exemples/dce_test

- dce_college_lot_unique : lot unique, RC au format Word, dates en chiffres,
  visite facultative, critère environnemental.
- dce_hopital_3_lots : 3 lots, formulations différentes du DCE Valbourg,
  visite obligatoire, qualification exigée, délai de questions relatif.
- dce_bailleur_incomplet : DCE volontairement incomplet (date limite absente,
  critères qui ne font pas 100 %, visite non mentionnée) pour vérifier que
  l'outil marque « à vérifier » au lieu d'inventer.

Tous les noms, adresses, montants et dates sont inventés.
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

import docx
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from generer_donnees_fictives import pdf  # noqa: E402


def _docx(chemin: Path, blocs: list[str]) -> None:
    d = docx.Document()
    for b in blocs:
        if b.startswith("# "):
            d.add_heading(b[2:], level=2)
        elif b != "---":
            for ligne in b.split("\n"):
                d.add_paragraph(ligne)
    d.save(chemin)


def _dpgf(chemin: Path, titre: str, postes: list[tuple[str, str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31]
    ws.append(["N°", "Désignation", "Unité", "Quantité", "PU HT", "Total HT"])
    for code, designation, unite in postes:
        ws.append([code, designation, unite, None, None, None])
    wb.save(chemin)


def _ae(chemin: Path, reference: str) -> None:
    d = docx.Document()
    d.add_heading("ACTE D'ENGAGEMENT", 0)
    d.add_paragraph(f"Marché n° {reference}. Le candidat s'engage sur le montant de son offre. Document daté et signé.")
    d.save(chemin)


# --- 1. Collège, lot unique -------------------------------------------------

RC_COLLEGE = [
    "# REGLEMENT DE LA CONSULTATION",
    "Département fictif de Loire-Bretagne - Direction des bâtiments\nN° de la consultation : CD-LB-2026-087",
    "# Article 1 - Objet",
    "Objet du marché : Mise en conformité électrique et remplacement de l'éclairage du collège Jules Verne\n"
    "Lieu d'exécution : 4 avenue des Marais, 44888 Saint-Jean-des-Prés\n"
    "Pouvoir adjudicateur : Département fictif de Loire-Bretagne, 3 quai Ceineray, 44000 Nantes\n"
    "Maîtrise d'oeuvre : Bureau d'études fictif Ohm Conseil, 17 rue Lamoricière, 44100 Nantes",
    "# Article 2 - Procédure",
    "Le marché est passé selon la procédure d'appel d'offres ouvert. Il n'est pas alloti : il fait l'objet d'un lot unique "
    "comprenant les courants forts et les courants faibles.",
    "# Article 3 - Visite",
    "Une visite de site facultative est proposée aux candidats le 24/11/2026 à 14h00, sur rendez-vous auprès du gestionnaire du collège.",
    "# Article 4 - Pièces à remettre",
    "Au titre de la candidature : DC1, DC2, attestations d'assurance en cours de validité, attestation de vigilance URSSAF, "
    "certificat de qualification Qualifelec ou équivalent, références de travaux similaires.\n"
    "Au titre de l'offre : acte d'engagement, DPGF, mémoire technique, planning prévisionnel.",
    "# Article 5 - Critères d'attribution",
    "1. Valeur technique : 50 %\n"
    "- Organisation du chantier pendant l'année scolaire : 25 %\n"
    "- Moyens humains et qualifications : 15 %\n"
    "- Qualité des matériels proposés : 10 %\n"
    "2. Prix : 40 %\n"
    "3. Performance environnementale : 10 %",
    "# Article 6 - Remise des offres",
    "Date limite de remise des offres : 15/12/2026 à 16h00.\n"
    "Les demandes de renseignements complémentaires sont adressées au plus tard le 04/12/2026 via la plateforme.\n"
    "Le délai de validité des offres est de 150 jours.\n"
    "Contact : batiments@loire-bretagne-fictif.fr",
]

CCAP_COLLEGE = [
    "# CAHIER DES CLAUSES ADMINISTRATIVES PARTICULIERES",
    "# Article 5 - Délais",
    "Le délai d'exécution est de 4 mois. Les travaux en salles de classe sont réalisés pendant les vacances d'hiver et de printemps.",
    "# Article 6 - Environnement",
    "Le titulaire assure le tri et la traçabilité des déchets de chantier. Les luminaires déposés sont confiés à un éco-organisme agréé. "
    "Le titulaire privilégie le covoiturage de ses équipes et limite les déplacements.",
    "# Article 7 - Insertion",
    "Le titulaire réserve au minimum 120 heures de travail à des personnes en parcours d'insertion.",
]

CCTP_COLLEGE = [
    "# CAHIER DES CLAUSES TECHNIQUES PARTICULIERES - LOT UNIQUE ELECTRICITE",
    "# Article 1 - Généralités",
    "Les installations respectent la norme NF C 15-100 et le règlement de sécurité contre l'incendie des ERP de type R. "
    "Le collège reste en activité : les coupures sont programmées le mercredi après-midi.",
    "# Article 2 - Eclairage",
    "Remplacement de 310 luminaires par des luminaires LED avec gradation automatique selon l'apport de lumière du jour.",
    "# Article 3 - Courants faibles",
    "Reprise du câblage informatique des salles de technologie en catégorie 6A. Interface avec le prestataire informatique du Département pour le brassage.",
]

# --- 2. Hôpital, 3 lots ------------------------------------------------------

RC_HOPITAL = [
    "# REGLEMENT DE CONSULTATION",
    "Centre hospitalier fictif de Saint-Aubin-sur-Erdre\nRéférence : CHSA-TX-2026-31",
    "# 1. Objet et allotissement",
    "La présente consultation porte sur la rénovation des installations électriques et de sécurité incendie du bâtiment B (médecine et SSR).\n"
    "Adresse du chantier : 2 rue de l'Hôpital, 44777 Saint-Aubin-sur-Erdre\n"
    "Maître d'ouvrage : Centre hospitalier fictif de Saint-Aubin-sur-Erdre, direction des travaux\n"
    "Maître d'oeuvre : Groupement fictif Arc & Fluides, 5 boulevard Guist'hau, 44000 Nantes\n"
    "Bureau d'études fluides : BET fictif Thermi-Conseil\n"
    "Contrôleur technique : Bureau fictif Contrôle Ouest\n"
    "Coordonnateur SPS : Cabinet fictif Prévention Chantiers 44\n"
    "OPC : Société fictive Planif'Bât",
    "Le marché est passé en procédure adaptée. Il est décomposé comme suit :",
    "Lot 1 : Courants forts - distribution et éclairage\nLot 2 : Système de sécurité incendie (SSI de catégorie A)\nLot 3 : Courants faibles - appel malade et contrôle d'accès",
    "Les candidats peuvent répondre à un ou plusieurs lots.",
    "# 2. Visite",
    "La visite du site est obligatoire. Elle se déroulera le 3 décembre 2026 à 9h30, rendez-vous à l'accueil du bâtiment B. "
    "Les offres des candidats n'ayant pas effectué la visite seront rejetées.",
    "# 3. Conditions de participation",
    "La qualification APSAD I7 et F7 est exigée pour le lot 2.\n"
    "Pièces de candidature : DC1, DC2, attestations d'assurance RC et décennale, attestation URSSAF, références de travaux en milieu hospitalier.\n"
    "Pièces de l'offre : acte d'engagement, BPU et DPGF du lot, mémoire technique, planning d'exécution.\n"
    "Le mémoire technique respecte obligatoirement le cadre de réponse joint en annexe du présent règlement.",
    "---",
    "# 4. Jugement des offres",
    "Les critères retenus sont :",
    "1. Valeur technique : 50 %\n"
    "- Méthodologie en établissement de santé occupé 24h/24 : 20 %\n"
    "- Continuité de la mise en sécurité incendie pendant les travaux : 15 %\n"
    "- Composition de l'équipe et habilitations : 15 %\n"
    "2. Prix : 40 %\n"
    "3. Délai d'exécution : 10 %",
    "# 5. Remise des plis",
    "Les plis devront parvenir par voie électronique avant le 8 janvier 2027 à 17h00.\n"
    "Les questions sont posées au plus tard 10 jours avant la date limite de remise des plis.\n"
    "Délai de validité des offres : 180 jours.\n"
    "Contact : travaux@ch-saint-aubin-fictif.fr",
]

CCAP_HOPITAL = [
    "# CCAP",
    "# Article 4 - Délais",
    "Le délai global d'exécution est de 8 mois. Les travaux dans les chambres sont réalisés par groupes de quatre chambres libérées par le service.",
    "# Article 5 - Hygiène",
    "Le titulaire respecte les consignes de l'équipe opérationnelle d'hygiène : confinement des zones, aspiration des poussières à la source, protocole de bionettoyage.",
    "# Article 6 - Développement durable",
    "Le titulaire limite les nuisances sonores entre 7h et 20h et trie les déchets de chantier. Les DEEE sont repris par une filière agréée.",
    "# Article 7 - Pénalités",
    "En cas de coupure non programmée de la sécurité incendie, une pénalité de 1 500 euros par heure est appliquée.",
]

CCTP_HOPITAL = {
    "1": ["# CCTP - LOT N°1 COURANTS FORTS",
          "# Article 1 - Généralités",
          "Les travaux se déroulent dans un établissement de santé en fonctionnement. Les circuits des chambres sont basculés sur une alimentation provisoire pendant les travaux.",
          "# Article 2 - Interfaces",
          "Le lot 1 alimente l'équipement de contrôle et de signalisation du lot 2 et les centrales d'appel malade du lot 3."],
    "2": ["# CCTP - LOT N°2 SYSTEME DE SECURITE INCENDIE",
          "# Article 1 - Objet",
          "Remplacement du SSI de catégorie A. L'ancien système reste en service jusqu'à la réception du nouveau.",
          "# Article 2 - Interfaces",
          "Les asservissements des portes coupe-feu sont réalisés en coordination avec le lot 3 (contrôle d'accès)."],
    "3": ["# CCTP - LOT N°3 COURANTS FAIBLES",
          "# Article 1 - Objet",
          "Remplacement de l'appel malade et extension du contrôle d'accès aux réserves de pharmacie.",
          "# Article 2 - Interfaces",
          "Le déverrouillage des accès sur alarme incendie est commandé par le SSI du lot 2."],
}

CADRE_HOPITAL = [
    "# ANNEXE 1 - CADRE DE REPONSE DU MEMOIRE TECHNIQUE",
    "Le candidat présente son mémoire technique en suivant les chapitres ci-dessous, dans cet ordre. 10 pages maximum par lot.",
    "Chapitre 1 - Organisation du chantier en établissement de santé occupé\n"
    "1.1 Phasage par groupes de chambres\n"
    "1.2 Mesures d'hygiène et de confinement\n"
    "Chapitre 2 - Maintien de la sécurité incendie pendant les travaux\n"
    "Chapitre 3 - Équipe affectée et habilitations\n"
    "Chapitre 4 - Planning d'exécution détaillé",
]

# --- 3. Bailleur, DCE incomplet ---------------------------------------------

RC_BAILLEUR = [
    "# REGLEMENT DE LA CONSULTATION",
    "Habitat fictif Atlantique - Résidence Les Mouettes (64 logements)\nRéférence de la consultation : HFA-2026-112",
    "# Article 1 - Objet",
    "Objet de la consultation : Remplacement des tableaux électriques des logements et mise en sécurité des parties communes\n"
    "Maître d'ouvrage : Habitat fictif Atlantique, 9 rue de la Marne, 44600 Saint-Nazaire",
    "# Article 2 - Procédure",
    "Procédure adaptée ouverte. Les travaux sont réalisés en milieu occupé, logement par logement, avec information préalable des locataires.",
    "# Article 3 - Pièces à fournir",
    "Candidature : DC1, DC2, attestations d'assurance, références. Offre : acte d'engagement, DPGF, mémoire technique.",
    "# Article 4 - Critères",
    "1. Prix : 60 %\n2. Valeur technique : 30 %",
    "# Article 5 - Remise des offres",
    "La date limite sera précisée sur le profil d'acheteur. Contact : marches@habitat-fictif-atlantique.fr",
]

CCTP_BAILLEUR = [
    "# CCTP - ELECTRICITE LOGEMENTS",
    "# Article 1 - Logements",
    "Remplacement du tableau de chaque logement par un tableau conforme avec interrupteurs différentiels 30 mA. Coupure limitée à une demi-journée par logement.",
    "# Article 2 - Parties communes",
    "Remplacement de l'éclairage des cages d'escalier par des hublots LED à détection.",
]


def generer(racine: Path) -> list[Path]:
    racine.mkdir(parents=True, exist_ok=True)
    dossiers = []

    d = racine / "dce_college_lot_unique"
    d.mkdir(exist_ok=True)
    _docx(d / "RC_college_Jules_Verne.docx", RC_COLLEGE)
    pdf(d / "CCAP_college.pdf", CCAP_COLLEGE)
    pdf(d / "CCTP_electricite.pdf", CCTP_COLLEGE)
    _ae(d / "AE_CD-LB-2026-087.docx", "CD-LB-2026-087")
    _dpgf(d / "DPGF_lot_unique.xlsx", "DPGF", [("1.1", "Luminaire LED 600x600 gradable", "u"),
                                                ("1.2", "Prise RJ45 catégorie 6A", "u"),
                                                ("1.3", "Tableau divisionnaire", "u")])
    dossiers.append(d)

    d = racine / "dce_hopital_3_lots"
    d.mkdir(exist_ok=True)
    pdf(d / "00_RC.pdf", RC_HOPITAL)
    pdf(d / "01_CCAP.pdf", CCAP_HOPITAL)
    for lot, blocs in CCTP_HOPITAL.items():
        pdf(d / f"02_CCTP_Lot_0{lot}.pdf", blocs)
        _dpgf(d / f"03_DPGF_Lot_0{lot}.xlsx", f"DPGF lot {lot}", [(f"{lot}.1", "Poste à chiffrer", "ens")])
    _ae(d / "04_AE.docx", "CHSA-TX-2026-31")
    pdf(d / "06_Annexe_cadre_de_reponse_memoire.pdf", CADRE_HOPITAL)
    (d / "05_plan_batiment_B.dwg").write_bytes(b"fichier DWG fictif")
    dossiers.append(d)

    d = racine / "dce_bailleur_incomplet"
    d.mkdir(exist_ok=True)
    pdf(d / "RC.pdf", RC_BAILLEUR)
    pdf(d / "CCTP.pdf", CCTP_BAILLEUR)
    _dpgf(d / "DPGF.xlsx", "DPGF", [("1", "Tableau électrique logement T3", "u")])
    dossiers.append(d)

    zips = []
    for d in dossiers:
        zips.append(Path(shutil.make_archive(str(d), "zip", d)))
    return zips


if __name__ == "__main__":
    for z in generer(Path(sys.argv[1] if len(sys.argv) > 1 else "exemples/dce_test")):
        print(z)
